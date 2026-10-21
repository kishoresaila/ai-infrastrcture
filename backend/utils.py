"""Shared config, text chunking and the deterministic storage calculator."""
import math
import os
from pathlib import Path

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434")
QDRANT_URL = os.getenv("QDRANT_URL", "http://localhost:6333")
LLM_MODEL = os.getenv("LLM_MODEL", "qwen3:8b")
EMBED_MODEL = os.getenv("EMBED_MODEL", "nomic-embed-text")
COLLECTION = os.getenv("COLLECTION", "infra_docs")
DOCS_DIR = Path(os.getenv("DOCS_DIR", "documents"))
EMBED_DIM = 768  # nomic-embed-text output size

CATEGORIES = ["linux", "docker", "kubernetes", "storage", "devops"]
ALLOWED_EXT = {".md", ".txt", ".pdf", ".docx"}
MAX_UPLOAD_MB = 20


def extract_text(path: Path) -> str:
    """Read plain text out of .md/.txt/.pdf/.docx files."""
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n\n".join(parts)
    raise ValueError(f"Unsupported file type: {ext}")


def chunk_text(text: str, size: int = 800, overlap: int = 120) -> list[str]:
    """Split text into overlapping chunks, preferring paragraph boundaries."""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks, current = [], ""
    for p in paragraphs:
        if len(current) + len(p) + 2 <= size:
            current = f"{current}\n\n{p}" if current else p
        else:
            if current:
                chunks.append(current)
            # carry a tail of the previous chunk for context
            tail = current[-overlap:] if current and overlap else ""
            current = f"{tail}\n\n{p}" if tail else p
            while len(current) > size * 2:  # very long paragraph
                chunks.append(current[:size])
                current = current[size - overlap:]
    if current:
        chunks.append(current)
    return chunks


# ---------------------------------------------------------------- calculator
def storage_calculator(users: int, gb_per_user: float, raid: str = "RAID5",
                       disk_size_gb: float = 1000, growth_percent: float = 20,
                       hot_spares: int = 1) -> dict:
    """Deterministic RAID sizing. Never ask the LLM to do maths."""
    raid = raid.upper().replace(" ", "")
    required = users * gb_per_user * (1 + growth_percent / 100)
    data_disks = math.ceil(required / disk_size_gb)  # disks worth of usable space

    if raid == "RAID0":
        n, parity, min_n = data_disks, 0, 2
    elif raid == "RAID1":
        n, parity, min_n = data_disks * 2, 0, 2
    elif raid == "RAID5":
        n, parity, min_n = data_disks + 1, 1, 3
    elif raid == "RAID6":
        n, parity, min_n = data_disks + 2, 2, 4
    elif raid == "RAID10":
        n = data_disks * 2
        n += n % 2
        parity, min_n = 0, 4
    else:
        raise ValueError(f"Unsupported RAID level: {raid}")

    n = max(n, min_n)
    if raid in ("RAID1", "RAID10"):
        usable = (n / 2) * disk_size_gb
    elif raid == "RAID0":
        usable = n * disk_size_gb
    else:
        usable = (n - parity) * disk_size_gb

    return {
        "raid_level": raid,
        "users": users,
        "gb_per_user": gb_per_user,
        "raw_requirement_gb": round(users * gb_per_user, 2),
        "required_with_growth_gb": round(required, 2),
        "disk_size_gb": disk_size_gb,
        "disks_required": n,
        "hot_spares": hot_spares,
        "total_disks_with_spares": n + hot_spares,
        "raw_capacity_gb": round(n * disk_size_gb, 2),
        "usable_capacity_gb": round(usable, 2),
        "fault_tolerance": {
            "RAID0": "none", "RAID1": "1 disk per mirror",
            "RAID5": "1 disk", "RAID6": "2 disks",
            "RAID10": "1 disk per mirror pair",
        }[raid],
    }
